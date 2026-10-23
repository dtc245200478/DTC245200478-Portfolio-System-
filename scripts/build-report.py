"""Generate the Link Me academic report from verified repository evidence.
Run with the Codex bundled Python runtime (python-docx, Pillow).
"""
from pathlib import Path
import json,subprocess
from docx import Document
from docx.shared import Cm,Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9);sec.different_first_page_header_footer=True
for name,size in [('Normal',12),('Title',21),('Subtitle',14),('Heading 1',15),('Heading 2',12.5),('Caption',10)]:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.1
 if name.startswith('Heading'):
  st.font.bold=True;st.paragraph_format.space_before=Pt(8);st.paragraph_format.keep_with_next=True
 if name=='Caption':st.font.italic=True
 for border in st.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=f.add_run('Link Me  |  ');rr.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.core_properties.author='Trần Trung Dũng';doc.core_properties.title='Báo cáo triển khai Link Me';doc.core_properties.subject='Đề 13 Website Portfolio Giới thiệu Cá nhân'
def p(text,style=None):
 x=doc.add_paragraph(text,style);x.paragraph_format.widow_control=True
 if style is None:x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return x
def h(text):return doc.add_heading(text,2)
def page(title):
 x=doc.add_heading(title,1);x.paragraph_format.page_break_before=True
 return x
def code(text):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(7);x.paragraph_format.line_spacing=1
 for i,line in enumerate(text.splitlines()):
  rr=x.add_run(('\n' if i else '')+line);rr.font.name='Consolas';rr.font.size=Pt(8.5);rr.font.color.rgb=RGBColor(0,0,0)
 return x
def table(headers,rows,widths=None,size=10):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,text in enumerate(headers):t.rows[0].cells[i].text=text
 if widths is None:widths=[6.5/len(headers)]*len(headers)
 for data in rows:
  for c,v in zip(t.add_row().cells,data):c.text=str(v)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');b.append(edge)
   pr.append(b);m=OxmlElement('w:tcMar')
   for side,n in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
    a=OxmlElement('w:'+side);a.set(qn('w:w'),n);a.set(qn('w:type'),'dxa');m.append(a)
   pr.append(m)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.05
    for run in pp.runs:run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p('') .paragraph_format.space_after=Pt(0)
 return t
figure_no=0
def caption(text):
 global figure_no
 figure_no+=1
 x=p('Hình '+str(figure_no)+'  '+text,'Caption');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
 return x
def image(path,label,width=6.2,maxheight=4.5):
 file=ROOT/path;assert file.exists(),file
 iw,ih=Image.open(file).size;actual=min(width,maxheight*iw/ih)
 x=p('');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(4)
 pic=x.add_run().add_picture(str(file),width=Inches(actual));pic._inline.docPr.set('descr',label)
 x.paragraph_format.keep_with_next=True;caption(label)
 return x
# 1 Cover
p('BÁO CÁO BÀI THỰC HÀNH','Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('TRIỂN KHAI VÀ QUẢN TRỊ\nHỆ THỐNG PHẦN MỀM','Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('\n')
x=p('Link Me','Title');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
for run in x.runs:run.font.size=Pt(32)
p('ĐỀ 13\nWEBSITE PORTFOLIO VÀ GIỚI THIỆU CÁ NHÂN','Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('\n')
for text in ['Sinh viên thực hiện  Trần Trung Dũng','Mã số sinh viên  dtc245200478','Lớp  CNTT K23D','Giảng viên hướng dẫn  Vũ Việt Dũng']:
 x=p(text);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('\n')
p('Tháng 10 năm 2026').alignment=WD_ALIGN_PARAGRAPH.CENTER
# 2
page('1 Tổng quan và mục lục')
p('Link Me là website portfolio cá nhân có trang quản trị nội dung. Bài thực hành tập trung vào việc triển khai và quản trị một hệ thống web hoàn chỉnh: ứng dụng, cơ sở dữ liệu, reverse proxy, giám sát, log tập trung và các biện pháp gia cố bảo mật. Toàn bộ dịch vụ được đóng gói bằng Docker Compose để có thể chạy lại từ mã nguồn trên GitHub.')
p('Kết quả kiểm thử được ghi trong thư mục docs/evidence. Cả Docker Desktop và Docker Engine độc lập trong Ubuntu WSL2 đều đã chạy website, kết nối MySQL, thu thập metrics và truy vấn log. Nội dung mẫu về kỹ năng, kinh nghiệm và hai dự án ý tưởng được ghi nhãn minh họa, không phải thành tích thực tế của sinh viên.')
h('Mục lục')
for text in ['1 Tổng quan và mục lục  .................................................. 2','2 Đối chiếu yêu cầu và quản lý mã nguồn  .......................... 3','3 Kiến trúc triển khai ứng dụng  ......................................... 4','4 Thiết kế dữ liệu và quản lý phiên  .................................... 5','5 Giao diện portfolio công khai  .......................................... 6','6 Quản trị nội dung và API  ................................................ 7','7 Triển khai bằng Docker Desktop  .................................... 8','8 Triển khai độc lập trong WSL2  ........................................ 9','9 Kiến trúc giám sát và metrics  .......................................... 10','10 Dashboard và phân tích số liệu  ..................................... 11','11 Thu thập log và truy vấn LogQL  .................................... 12','12 Gia cố bảo mật hệ thống  ............................................... 13','13 Sao lưu phục hồi và kiểm thử  ........................................ 14','14 Demo và kết luận  .......................................................... 15']:
 x=p(text);x.paragraph_format.space_after=Pt(3)
# 3
page('2 Đối chiếu yêu cầu và quản lý mã nguồn')
p('Đề 13 yêu cầu website giới thiệu cá nhân có admin quản lý nội dung, sử dụng MySQL và phpMyAdmin. Những yêu cầu chung tại trang 1 và 2 của đề được triển khai đầy đủ theo bảng đối chiếu sau. Giao diện của Link Me được thiết kế riêng vì đề không cung cấp mẫu giao diện bắt buộc.')
table(['Yêu cầu','Kết quả trong Link Me'],[
('GitHub và README','Repository theo MSSV; README tiếng Việt, source, cấu hình, script và minh chứng.'),
('Ứng dụng và DB','Next.js, MySQL 8.4, migration/seed và phpMyAdmin hoạt động.'),
('Nginx reverse proxy','Website qua Nginx; security headers, giới hạn request/upload và timeout.'),
('Prometheus và Grafana','5 target, số liệu container/web/DB, dashboard tự provision.'),
('Loki và Promtail','Log app, Nginx, MySQL; 3 truy vấn LogQL có kết quả.'),
('Hardening','Non-root, phân mạng, secret mạnh, giảm quyền, chống CSRF và giới hạn đăng nhập.'),
('Báo cáo và demo','Báo cáo Word/PDF, sơ đồ có nguồn chỉnh sửa, ảnh chụp và kịch bản demo.')],[1.7,4.7],10.5)
h('Quản lý phiên bản')
p('Repository sử dụng tài khoản GitHub dtc245200478 và giữ commit khởi tạo. Các thay đổi được commit theo ba mốc có nội dung tương ứng yêu cầu đề. Tên tác giả được cấu hình cục bộ cho repository để không sử dụng nhầm danh tính Git của máy.')
table(['Mốc','Commit và nội dung'],[('1','ea69d04  Website/admin, MySQL, phpMyAdmin và Nginx.'),('2','d6a9c41  Prometheus, Grafana và metrics container/web/DB.'),('3','Logging, hardening, sao lưu và tài liệu bàn giao. Báo cáo này thuộc chính commit mốc 3; SHA cuối được xem trong git log trên GitHub.')],[.65,5.75],10.5)
code('https://github.com/dtc245200478/\nDTC245200478-Portfolio-System-\ngit log --oneline -4')
p('Mật khẩu, phiên chạy, file .env, backup và kết quả tạm nằm ngoài Git. Các gói npm và image được khóa phiên bản để hạn chế thay đổi ngoài dự kiến khi dựng lại bài.')
# 4
page('3 Kiến trúc triển khai ứng dụng')
p('Trình duyệt truy cập cổng 8088 tại localhost. Nginx tiếp nhận request và chuyển đến Next.js trong mạng Docker. Ứng dụng xử lý HTML, API quản trị, xác thực và truy vấn MySQL. MySQL không publish cổng ra host. Công cụ phpMyAdmin có cổng localhost riêng để đăng nhập quản trị database.')
image('docs/diagrams/application.png','Luồng triển khai ứng dụng từ Compose và cấu hình Nginx',width=6.1,maxheight=5.35)
p('Dữ liệu nội dung nằm trong volume mysql_data; ảnh upload nằm trong uploads. App ghi log JSON vào app_logs. Các volume này có vòng đời độc lập với container nên restart hoặc tạo lại container không xóa dữ liệu. Container migration chạy sau khi MySQL sẵn sàng qua TCP; app chỉ khởi động khi migration hoàn tất.')
table(['Mạng','Vai trò'],[('frontend','Nginx liên lạc với ứng dụng.'),('database','App, migration, MySQL, phpMyAdmin và DB exporter.'),('monitoring','App, exporter, Prometheus, Grafana, Loki và Promtail.'),('edge và tools','Cổng localhost của reverse proxy và công cụ quản trị.')],[1.3,5.1],10)
# 5
page('4 Thiết kế dữ liệu và quản lý phiên')
p('Database linkme dùng InnoDB, utf8mb4 và MySQL 8.4. Ứng dụng gửi giá trị bằng placeholder để tránh ghép dữ liệu đầu vào vào SQL. Tên bảng và cột chỉ được chọn từ allowlist trên server. Schema gồm 11 bảng, chia theo mục đích sử dụng dưới đây.')
table(['Bảng','Nội dung chính','Ràng buộc chính'],[
('admins','Tài khoản và hash mật khẩu scrypt','PK id; username UNIQUE'),('sessions','Hash token, admin_id, hạn phiên','PK token_hash; FK admins'),('login_attempts','Số lần thử và thời điểm cửa sổ','PK attempt_key'),('profile','Hồ sơ, email, ảnh, lời mời liên hệ','Một bản ghi id 1'),('skills','Tên, nhóm, mức độ, thứ tự','level từ 0 đến 100'),('education','Ngành học, lớp, thời gian, mô tả','PK id'),('experience','Vai trò, đơn vị, thời gian, mô tả','PK id'),('projects','Slug, nội dung, ảnh, link, công khai','slug UNIQUE'),('social_links','Nhãn, URL và thứ tự','PK id'),('media','Tên ảnh UUID, MIME, dung lượng','filename UNIQUE'),('schema_migrations','Phiên bản seed/migration đã áp dụng','PK version')],[1.1,3.4,1.9],9.5)
image('docs/diagrams/auth.png','Phạm vi xác thực với quan hệ tài khoản và phiên',width=3.25,maxheight=3.25)
p('Quan hệ admins đến sessions là 1 đến 0..n; admin_id không được NULL và xóa tài khoản sẽ xóa phiên liên quan. Các bảng nội dung được quản lý trực tiếp bởi một admin. media được tham chiếu bằng đường dẫn ảnh, không có FK từ projects/profile đến media. Seed chỉ chạy một lần theo schema_migrations và không đặt lại mật khẩu đã đổi.')
# 6
page('5 Giao diện portfolio công khai')
p('Giao diện Link Me sử dụng nền sáng và màu xanh làm điểm nhấn. Trang chủ dẫn người xem từ phần giới thiệu đến kỹ năng, dự án, học vấn, kinh nghiệm và liên hệ. Các liên kết nội bộ cho phép di chuyển trực tiếp đến từng phần; dự án có trang chi tiết riêng và nút xem mã nguồn khi có URL hợp lệ.')
x=p('');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.keep_with_next=True
for path,width in [('docs/evidence/website-hero.png',4.55),('docs/evidence/website-mobile-viewport.png',1.65)]:
 pic=x.add_run().add_picture(str(ROOT/path),width=Inches(width));pic._inline.docPr.set('descr','Giao diện Link Me '+path)
caption('Ảnh thật của website ở desktop và chiều rộng 390 pixel')
h('Hiển thị nội dung và khả năng sử dụng')
p('Nội dung được đọc động từ MySQL, vì vậy thay đổi trong admin xuất hiện sau khi lưu và tải lại trang. Dự án chưa công khai không có trong danh sách và truy cập trực tiếp URL trả 404. Slug là duy nhất để đường dẫn chi tiết ổn định. React tự escape nội dung văn bản; server không cho nhập HTML tùy ý.')
p('Bố cục chuyển từ hai cột sang một cột trên màn hình nhỏ. Ảnh dùng SVG minh họa có sẵn hoặc WebP do quản trị tải lên. Các nút có trạng thái focus, label và văn bản rõ ràng; form phản hồi lỗi ngay trong giao diện. Kiểm tra responsive đã sửa phần vòng trang trí gây tràn và xác nhận chiều rộng nội dung bằng chiều rộng viewport 390 pixel.')
p('Họ tên, MSSV và lớp sử dụng thông tin đã cung cấp. Phần kỹ năng, kinh nghiệm, ảnh đại diện và dự án ý tưởng có nhãn minh họa. Người quản trị có thể thay bằng nội dung cá nhân thật sau khi bàn giao.')
# 7
page('6 Quản trị nội dung và API')
p('Trang /admin phục vụ một quản trị viên. Sau khi đăng nhập, người dùng sửa hồ sơ, thêm/sửa/xóa kỹ năng, học vấn, kinh nghiệm, dự án và liên kết. Dự án có trạng thái bản nháp hoặc công khai và trường sort_order để sắp xếp. Thư viện ảnh lưu ảnh đã qua kiểm tra và cung cấp URL để chọn làm ảnh đại diện hoặc ảnh dự án.')
image('docs/evidence/admin-profile.png','Trang quản trị hồ sơ sau khi lưu thành công',width=6.1,maxheight=4.15)
h('Xác thực và kiểm tra đầu vào')
p('Mật khẩu được hash bằng scrypt với salt riêng. Server sinh token ngẫu nhiên 32 byte, chỉ lưu SHA-256 của token trong MySQL và đặt cookie HttpOnly, SameSite Strict, thời hạn 8 giờ. CSRF token được suy ra bằng HMAC từ token phiên và secret; mọi thao tác ghi kiểm tra Origin và CSRF. Đổi mật khẩu thu hồi toàn bộ phiên để yêu cầu đăng nhập lại.')
p('Các API /api/admin/* kiểm tra quyền trước khi đọc dữ liệu quản trị. POST/PUT dùng schema allowlist; DELETE không được xóa hồ sơ đơn. Trùng slug trả 409, thiếu dữ liệu trả 400, chưa đăng nhập trả 401 và CSRF sai trả 403. Upload chỉ nhận JPG/PNG/WebP không quá 5 MB; sharp giải mã, giới hạn số pixel, thay tên UUID và xuất WebP để tránh lưu nội dung thực thi.')
# 8
page('7 Triển khai bằng Docker Desktop')
p('Môi trường Windows dùng Docker Desktop với Linux containers và context desktop-linux. Máy kiểm thử có Docker Engine Desktop 29.7.2. Repository được clone vào một thư mục Git riêng; script tạo .env và secret một lần trước khi build. Người chạy không cần cài Node hoặc MySQL trên Windows.')
code('git clone https://github.com/dtc245200478/\nDTC245200478-Portfolio-System-.git link-me\ncd link-me\npowershell -ExecutionPolicy Bypass -File scripts/start.ps1\npowershell -ExecutionPolicy Bypass -File scripts/test.ps1')
p('Khi copy lệnh clone, ghép hai dòng URL thành một dòng. README trong repository có lệnh đầy đủ để chạy trực tiếp.')
table(['Thành phần','Địa chỉ mặc định','Tài khoản'],[('Website và admin','localhost:8088 và /admin','admin'),('phpMyAdmin','localhost:8089','linkme_admin'),('Grafana','localhost:3004','admin'),('Prometheus','localhost:9093','Chỉ localhost')],[1.5,3.1,1.8],10)
image('docs/evidence/phpmyadmin.png','phpMyAdmin đăng nhập thành công và liệt kê bảng linkme',width=6.05,maxheight=3.6)
p('Mật khẩu tương ứng nằm trong .local/secrets; không đưa vào báo cáo hoặc Git. Script kiểm tra engine và cổng, build image, đợi healthcheck. MySQL phải sẵn sàng qua TCP để tránh nhầm với server tạm lúc khởi tạo. Dừng bằng scripts/stop.ps1 giữ nguyên named volume.')

# 9
page('8 Triển khai độc lập trong WSL2')
p('Môi trường thứ hai là Ubuntu 24.04.5 LTS trong WSL2, Docker Engine 29.9.1 và Compose 5.6.0. Ubuntu đã có Docker Engine riêng; script Bash đặt DOCKER_HOST=unix:///var/run/docker.sock và bỏ DOCKER_CONTEXT để tránh dùng nhầm engine của Docker Desktop. Mã nguồn nên nằm trong filesystem Linux để tăng tốc build và tránh khác biệt quyền file từ NTFS.')
code('git clone https://github.com/dtc245200478/\nDTC245200478-Portfolio-System-.git ~/link-me\ncd ~/link-me\nsh scripts/start.sh\nsh scripts/test.sh\ndocker compose ps\ndocker compose exec -T app node --input-type=module \
  < scripts/check-services.mjs\ndocker compose exec -T app node --input-type=module \
  < scripts/verify-logs.mjs')
h('Phân biệt hai engine')
table(['Thuộc tính','Docker Desktop','Ubuntu WSL Engine'],[('Lựa chọn engine','context desktop-linux','socket Unix của Ubuntu'),('Cổng website kiểm thử','8088','8188'),('phpMyAdmin kiểm thử','8089','8189'),('Grafana và Prometheus','3004 và 9093','3104 và 9193'),('Dữ liệu volume','Engine của Desktop','Engine của Ubuntu'),('Desktop tắt','Dịch vụ Desktop dừng','Dịch vụ WSL tiếp tục hoạt động')],[1.7,2.35,2.35],10.5)
p('Các cổng khác nhau ở bài kiểm thử cho phép chạy hai bản đồng thời. Khi chỉ chạy một engine, dùng cổng mặc định 8088/8089/3004/9093. .env cho phép thay cổng; không cần sửa source. Hai engine không tự dùng chung volume. Để chuyển bài sang máy hoặc engine khác, dùng backup gồm SQL và ảnh, kiểm tra checksum rồi phục hồi.')
p('Kiểm thử WSL dùng socket Unix và Docker Engine riêng của Ubuntu; kiểm thử Desktop dùng context desktop-linux. Hai daemon được kiểm tra độc lập và dùng cổng khác nhau. Những container của bài đề 16 sẵn có trong Ubuntu được giữ nguyên. Không cài chồng Desktop WSL Integration lên distro có engine riêng.')
h('Tình huống vận hành')
p('Nếu Docker chưa chạy, kiểm tra docker info và dịch vụ Docker của Ubuntu. Nếu cổng bận, đổi .env trước khi chạy lại. Khi MySQL đã có dữ liệu, không đổi secret bằng cách sửa file mật khẩu ban đầu. Scripts/stop.sh chỉ dừng project; không xóa volume hoặc prune tài nguyên toàn máy.')
# 10
page('9 Kiến trúc giám sát và metrics')
p('Prometheus scrape mỗi 15 giây. App xuất metrics nội bộ qua /metrics; Nginx exporter đọc stub_status; MySQL exporter dùng tài khoản riêng; cAdvisor đọc cgroup và runtime của Docker host. Nhãn project/service giúp lọc đúng các container thuộc Link Me, tránh nhầm số liệu bài đề 16 đang chạy cùng máy.')
image('docs/diagrams/observability.png','Luồng scrape metrics và đẩy log trong Link Me',width=6.2,maxheight=3.7)
table(['Job','Đích scrape','Ý nghĩa'],[('prometheus','prometheus:9090','Sức khỏe Prometheus'),('linkme','app:3000/metrics','API, độ trễ, lượt xem'),('nginx','nginx-exporter:9113','Lưu lượng và kết nối'),('mysql','mysqld-exporter:9104','Kết nối và truy vấn DB'),('cadvisor','cadvisor:8080','CPU/RAM container')],[1.25,2.5,2.65],10)
p('Hai engine đều được kiểm tra target up và có chuỗi CPU/RAM của app, Nginx và MySQL. Chỉ thấy số liệu CPU tổng của host không đủ để nghiệm thu. Cảnh báo LinkMeTargetDown kích hoạt khi một target mất trên một phút; bài thực hành chưa cấu hình đầu nhận thông báo bên ngoài.')
p('Giữ dữ liệu Prometheus 7 ngày, giới hạn dung lượng 512 MB. Bộ đếm ứng dụng bắt đầu lại khi app restart; dữ liệu portfolio trong MySQL vẫn được giữ. Cần chờ ít nhất hai lần scrape để biểu đồ rate thể hiện có ý nghĩa.')
# 11
page('10 Dashboard và phân tích số liệu')
p('Grafana tự nạp datasource Prometheus/Loki và dashboard Link Me Tổng quan vận hành. Bản bàn giao có 13 panel: 10 panel metrics và 3 panel log. Không cần import dashboard thủ công sau khi clone. Tài khoản admin có mật khẩu sinh riêng, không bật truy cập ẩn danh hoặc đăng ký công khai.')
image('docs/evidence/grafana-containers.png','Grafana hiển thị CPU RAM và lưu lượng Nginx thực tế',width=6.25,maxheight=4.3)
h('Đọc các nhóm số liệu')
p('CPU được tính từ tốc độ tăng container_cpu_usage_seconds_total, nhân 100 để biểu diễn phần trăm một lõi CPU; RAM sử dụng container_memory_working_set_bytes. Dùng các nhãn Compose project/service để nhóm theo từng dịch vụ. MySQL có panel số kết nối và tốc độ truy vấn; Nginx có panel request mỗi giây và kết nối đang hoạt động.')
code('sum by(container_label_com_docker_compose_service)(\n  rate(container_cpu_usage_seconds_total{\n    container_label_com_docker_compose_project="linkme13",\n    image!=""}[1m])) * 100\nmysql_global_status_threads_connected\nrate(nginx_http_requests_total[1m])')
p('Panel API theo status giúp phân biệt phản hồi thành công, sai thông tin và bị chặn. Độ trễ p95 được tính bằng histogram_quantile trên histogram API. Ảnh dashboard và kết quả scripts/check-services.mjs là minh chứng độc lập: một bên cho thấy giao diện, một bên xác nhận các chuỗi và target đúng tên dịch vụ.')
# 12
page('11 Thu thập log và truy vấn LogQL')
p('App ghi JSON theo time, level và event. Nginx ghi JSON có method, path, status, duration và bytes; đường dẫn không chứa query string. MySQL ghi lỗi vào volume riêng. Promtail đọc các volume ở chế độ chỉ đọc, lưu positions để tiếp tục sau restart và đẩy log đến Loki qua mạng monitoring.')
image('docs/evidence/logql-login.png','Grafana Explore trả các log đăng nhập thất bại',width=6.2,maxheight=3.6)
h('Ba truy vấn LogQL đã kiểm chứng')
code('{project="linkme13",service="nginx",stream="access"} | json\n\n{project="linkme13",service="app"} | json\n  | event="login_failed"\n\n{project="linkme13",service="app"} | json | level="error"')
p('Truy vấn thứ nhất phục vụ xem request web; truy vấn thứ hai tìm lần đăng nhập sai; truy vấn thứ ba tìm lỗi ứng dụng. Bộ kiểm thử tạo một lỗi có kiểm soát qua API diagnostics có xác thực và CSRF để minh chứng, không phải lỗi vận hành chưa xử lý. Log không ghi mật khẩu, cookie hay token phiên.')
p('Logrotate kiểm tra mỗi phút, xoay file khi vượt 5 MB, giữ 3 bản; Loki giữ 7 ngày. stdout của container cũng được giới hạn bằng json-file 5 MB và 3 file. Promtail đã hết hỗ trợ ngày 02/03/2026, được giữ theo yêu cầu môn học; khi vận hành dài hạn cần lập kế hoạch thay bằng Grafana Alloy.')
# 13
page('12 Gia cố bảo mật hệ thống')
p('Hardening được áp dụng đồng thời ở container, mạng, database, ứng dụng và reverse proxy. Kết quả kiểm tra Docker inspect cho cả hai engine được lưu trong docs/evidence/hardening.json. Bảng sau liên kết biện pháp với cấu hình và điều kiện kiểm chứng.')
table(['Biện pháp','Cấu hình áp dụng','Kiểm chứng'],[
('Non-root','App UID 1000; Nginx UID 101; MySQL UID 999','Docker inspect Config.User'),
('Giảm quyền','cap_drop ALL và no-new-privileges cho app/Nginx/MySQL','Inspect và Compose'),
('Filesystem chỉ đọc','Root filesystem read_only; tmpfs/vùng ghi riêng','Inspect và upload thực tế'),
('Cách ly mạng','frontend/database/monitoring nội bộ','Không publish MySQL, Loki, exporter'),
('Cổng localhost','Nginx, DB tool, Grafana, Prometheus bind 127.0.0.1','Inspect PortBindings'),
('Quyền database','App chỉ DML; migration DDL; exporter chỉ giám sát','GRANT trong init script'),
('Secret và mật khẩu','Sinh ngẫu nhiên, giữ ngoài Git; scrypt và salt','Quét staged file và DB hash'),
('Xác thực web','HttpOnly, SameSite Strict, Origin và CSRF','401/403 trong kiểm thử'),
('Chống dò mật khẩu','DB giới hạn 5 lần trong 15 phút; Nginx rate limit','Chuỗi 401 × 5 rồi 429'),
('Nginx và upload','CSP, nosniff, chống nhúng, giới hạn dung lượng','HTTP headers, từ chối file sai')],[1.25,3.1,2.05],10)
h('Ngoại lệ quyền được kiểm soát')
p('cAdvisor cần privileged và đọc cgroup/socket của host để thu thập số liệu container. storage-init và log-storage-init chạy root một lần để đặt quyền volume. logrotate dùng root với capability cần thiết để xoay các file khác UID. phpMyAdmin Apache khởi tạo bằng root, worker phục vụ chạy www-data. Vì vậy báo cáo không khẳng định mọi container đều non-root.')
p('Bản demo dùng HTTP localhost với security headers theo lựa chọn được đề cho phép. Cookie Secure chưa bật cho HTTP. Nếu triển khai Internet, phải bổ sung HTTPS, bật COOKIE_SECURE, rà soát origin, tài khoản, firewall và thay Promtail. CSP còn unsafe-inline do cơ chế hydration của Next.js; không coi cấu hình localhost là cấu hình production Internet hoàn chỉnh.')
# 14
page('13 Sao lưu phục hồi và kiểm thử')
p('Backup gồm database.sql, uploads.tar.gz và manifest.json có SHA-256, đặt trong .local/backups. Hai script cùng định dạng để có thể chuyển giữa Windows và WSL. Dữ liệu backup có hash tài khoản và phiên nên không commit. Restore kiểm tra tên project và checksum, yêu cầu xác nhận, dừng app/Nginx rồi khôi phục database và ảnh.')
code('powershell -File scripts/backup.ps1\npowershell -File scripts/restore.ps1 -BackupDir <backup>\n\nsh scripts/backup.sh\nsh scripts/restore.sh <backup>')
p('Kiểm thử cố ý đổi lời mời liên hệ và ghi đè một ảnh thử, sau đó phục hồi backup. Hash của các bảng nội dung, tài khoản, phiên và file ảnh khớp trước/sau phục hồi trên từng engine. File ảnh cũ không còn tham chiếu được giữ trên volume. Mật khẩu sau restore đi theo database đã sao lưu, không đi theo file mật khẩu khởi tạo của máy đích.')
table(['Nhóm kiểm thử','Kết quả thực tế'],[('Trang công khai và 404','Trang chủ/chi tiết hoạt động; bản nháp và slug sai trả 404.'),('Đăng nhập và cookie','Đúng/sai mật khẩu, phiên, HttpOnly, SameSite Strict.'),('Phân quyền và CSRF','Chưa đăng nhập trả 401; Origin/CSRF sai trả 403.'),('CRUD nội dung','Hồ sơ, kỹ năng, học vấn, kinh nghiệm, dự án và link.'),('Dữ liệu sai và trùng','Mức kỹ năng sai trả 400; slug trùng trả 409.'),('Upload','Ảnh hợp lệ thành WebP; SVG/file lỗi bị từ chối.'),('Đổi mật khẩu','Đổi thành công, thu hồi phiên, phục hồi mật khẩu kiểm thử.'),('Restart và seed','Hash dữ liệu/ảnh không đổi; seed không nhân đôi.'),('Giám sát và log','5 target up, CPU/RAM đúng dịch vụ, 3 LogQL có kết quả.'),('Sao lưu và phục hồi','SQL và ảnh phục hồi đúng sau thay đổi cố ý.')],[1.8,4.6],10)
p('Bộ smoke test đạt 53 kiểm tra HTTP trên mỗi engine. Các kiểm tra bổ sung bao gồm giới hạn đăng nhập, kiểm tra quyền container, giám sát, log, restart/seed, backup/restore và responsive. Các file JSON minh chứng ghi rõ thời điểm và kết quả; không thay thế lỗi bằng ảnh minh họa hoặc kết quả giả định.')
# 15
page('14 Demo và kết luận')
h('Kịch bản demo đề xuất')
for text in [
'1. Mở repository, README và lịch sử ba mốc triển khai trên GitHub.',
'2. Chạy script khởi động và xem trạng thái container của project linkme13.',
'3. Truy cập portfolio qua Nginx tại localhost:8088 và mở một dự án.',
'4. Đăng nhập admin; sửa hồ sơ và kiểm tra nội dung công khai cập nhật.',
'5. Thêm dự án bản nháp, công khai dự án rồi mở trang chi tiết.',
'6. Upload ảnh hợp lệ; thử file SVG để thấy cơ chế từ chối.',
'7. Mở phpMyAdmin, đăng nhập và xem các bảng của database linkme.',
'8. Mở Prometheus targets và dashboard Grafana về CPU/RAM, web, DB.',
'9. Đăng nhập sai một lần và chạy ba truy vấn LogQL trong Explore.',
'10. Giải thích non-root, phân mạng, secret, quyền DB và ngoại lệ cAdvisor.',
'11. Chạy backup/restore trên dữ liệu kiểm thử và đối chiếu kết quả.',
'12. Dừng Desktop và chứng minh hệ thống chạy bằng Engine của WSL.']:
 x=p(text);x.paragraph_format.space_after=Pt(4)
h('Kết luận')
p('Link Me đáp ứng chức năng portfolio có quản trị nội dung và yêu cầu triển khai của đề 13. Hệ thống kết hợp Docker Compose, MySQL, phpMyAdmin, Nginx, Prometheus, Grafana, Loki và Promtail; dữ liệu được giữ qua restart, có sao lưu/phục hồi và các lớp kiểm soát bảo mật phù hợp demo localhost. Hai cách chạy được kiểm thử riêng để tránh phụ thuộc nhầm vào Docker Desktop.')
p('Hướng phát triển tiếp theo là đưa nội dung cá nhân thật vào portfolio, bổ sung HTTPS cho môi trường có tên miền, dùng hệ thống gửi cảnh báo và chuyển từ Promtail sang tác nhân được hỗ trợ. Những phần mở rộng này không làm thay đổi kết quả nghiệm thu của bài thực hành hiện tại.')
h('Tài liệu tham khảo')
for text in [
'1. Đề tài thực hành môn Triển khai và Quản trị Hệ thống Phần mềm K23, trang 1, 2 và 5.',
'2. Next.js Self Hosting: https://nextjs.org/docs/app/guides/self-hosting',
'3. Docker Desktop WSL2: https://docs.docker.com/desktop/features/wsl/',
'4. Docker Engine Ubuntu: https://docs.docker.com/engine/install/ubuntu/',
'5. Grafana Promtail và thông báo EOL: https://grafana.com/docs/loki/latest/send-data/promtail/',
'6. Mã nguồn và kết quả kiểm thử: repository GitHub dtc245200478/DTC245200478-Portfolio-System-.']:
 x=p(text);x.paragraph_format.space_after=Pt(3)
 for run in x.runs:run.font.size=Pt(10)
output=ROOT/'docs/Bao_Cao_Link_Me_Tran_Trung_Dung.docx';doc.save(output)
print(output)
